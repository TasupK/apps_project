from __future__ import annotations

import io
import json
import os
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st

from agents.evaluation_agent import evaluate_candidates_from_phase2_results
from agents.monitor_agent import scan_inventory
from agents.reporting_agent import write_evaluation_report
from agents.web_research.agent import run_web_research
from db.database import get_connection
from graph.state import make_initial_state
from graph.workflow import build_graph


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
PO_FILE = BASE_DIR / "PO_Result.csv"
ENV_FILE = BASE_DIR / ".env"
BUDGET_LIMIT_KRW = 5_000_000


st.set_page_config(
    page_title="BuyBee",
    layout="wide",
    page_icon="🐝",
    initial_sidebar_state="expanded",
)


PROFESSIONAL_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Syne:wght@400;600;700;800&family=DM+Mono:wght@300;400;500&display=swap');

/* ── BuyBee · Clean Light Theme ─────────────────────────────── */
:root {
    --bg:           #F8F9FC;
    --surface:      #FFFFFF;
    --surface-2:    #F3F5F9;
    --surface-3:    #EDF0F6;
    --border:       #E4E8F0;
    --border-light: #D0D7E6;
    --accent:       #F59E0B;
    --accent-dim:   rgba(245,158,11,0.08);
    --accent-glow:  rgba(245,158,11,0.18);
    --navy:         #1E3A5F;
    --blue:         #2563EB;
    --blue-dim:     rgba(37,99,235,0.08);
    --success:      #059669;
    --success-dim:  rgba(5,150,105,0.08);
    --danger:       #DC2626;
    --danger-dim:   rgba(220,38,38,0.08);
    --text:         #374151;
    --text-bright:  #111827;
    --text-muted:   #6B7280;
    --text-dim:     #9CA3AF;
}

/* ── Streamlit chrome 제거 ──────────────────────────────────── */
#MainMenu { visibility: hidden; }
footer    { visibility: hidden; }
header    { visibility: hidden; height: 0; }

/* ── Base ───────────────────────────────────────────────────── */
.stApp {
    background-color: var(--bg) !important;
    font-family: 'DM Mono', monospace;
    color: var(--text);
}
.block-container {
    padding-top: 1.5rem !important;
    padding-bottom: 2rem !important;
    max-width: 1440px !important;
}

/* ── Sidebar ────────────────────────────────────────────────── */
[data-testid="stSidebar"] {
    background: var(--navy) !important;
    border-right: none !important;
}
[data-testid="stSidebar"] * { color: #BFD0E8 !important; font-family: 'DM Mono', monospace !important; }
[data-testid="stSidebar"] [data-baseweb="select"] > div {
    background: rgba(255,255,255,0.07) !important;
    border: 1px solid rgba(255,255,255,0.12) !important;
    border-radius: 6px !important;
    color: #F1F5F9 !important;
}
[data-testid="stSidebar"] .stButton > button {
    background: var(--accent-dim) !important;
    border: 1px solid rgba(245,158,11,0.5) !important;
    color: var(--accent) !important;
}
[data-testid="stSidebar"] [data-testid="stAlert"] {
    background: rgba(245,158,11,0.1) !important;
    border: 1px solid rgba(245,158,11,0.25) !important;
    border-radius: 6px !important;
    color: #FCD34D !important;
}

/* ── Metric Cards ───────────────────────────────────────────── */
[data-testid="stMetric"] {
    background: var(--surface) !important;
    border: 1px solid var(--border) !important;
    border-top: 3px solid var(--accent) !important;
    border-radius: 10px !important;
    padding: 1rem 1.25rem !important;
    box-shadow: 0 1px 4px rgba(0,0,0,0.06) !important;
}
[data-testid="stMetric"] label {
    color: var(--text-muted) !important;
    font-size: 0.62rem !important;
    font-weight: 500 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.12em !important;
    font-family: 'DM Mono', monospace !important;
}
[data-testid="stMetricValue"] {
    color: var(--text-bright) !important;
    font-size: 1.4rem !important;
    font-weight: 700 !important;
    font-family: 'Syne', sans-serif !important;
}

/* ── Buttons ────────────────────────────────────────────────── */
.stButton > button {
    background: var(--navy) !important;
    color: #fff !important;
    border: none !important;
    border-radius: 6px !important;
    font-weight: 600 !important;
    font-size: 0.78rem !important;
    font-family: 'DM Mono', monospace !important;
    letter-spacing: 0.06em !important;
    text-transform: uppercase !important;
    padding: 0.5rem 1.25rem !important;
    box-shadow: 0 1px 3px rgba(30,58,95,0.25) !important;
    transition: all 0.15s ease !important;
}
.stButton > button:hover {
    background: #162d4a !important;
    box-shadow: 0 4px 12px rgba(30,58,95,0.35) !important;
    transform: translateY(-1px) !important;
}

/* ── DataFrames ─────────────────────────────────────────────── */
[data-testid="stDataFrame"] {
    border: 1px solid var(--border) !important;
    border-radius: 8px !important;
    overflow: hidden !important;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04) !important;
}

/* ── Alerts ─────────────────────────────────────────────────── */
[data-testid="stAlert"] {
    border-radius: 8px !important;
    font-size: 0.82rem !important;
    font-family: 'DM Mono', monospace !important;
}

/* ── Chat ───────────────────────────────────────────────────── */
[data-testid="stChatMessage"] {
    background: var(--surface) !important;
    border: 1px solid var(--border) !important;
    border-radius: 8px !important;
    margin-bottom: 6px !important;
    font-family: 'DM Mono', monospace !important;
    font-size: 0.82rem !important;
}
[data-testid="stChatInput"] textarea {
    background: var(--surface) !important;
    border: 1px solid var(--border) !important;
    border-radius: 8px !important;
    color: var(--text-bright) !important;
    font-family: 'DM Mono', monospace !important;
    font-size: 0.82rem !important;
}
[data-testid="stChatInput"] textarea:focus {
    border-color: var(--accent) !important;
    box-shadow: 0 0 0 2px var(--accent-glow) !important;
}

/* ── Typography ─────────────────────────────────────────────── */
h1, h2, h3 {
    font-family: 'Syne', sans-serif !important;
    color: var(--text-bright) !important;
    letter-spacing: -0.02em !important;
}
.stCaption { color: var(--text-muted) !important; font-size: 0.72rem !important; }
hr { border-color: var(--border) !important; margin: 1.25rem 0 !important; }

/* ── Charts ─────────────────────────────────────────────────── */
[data-testid="stVegaLiteChart"],
[data-testid="stArrowVegaLiteChart"] {
    background: var(--surface) !important;
    border: 1px solid var(--border) !important;
    border-radius: 10px !important;
    padding: 1rem !important;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04) !important;
}

/* ── Selectbox ──────────────────────────────────────────────── */
[data-baseweb="select"] > div {
    background: var(--surface) !important;
    border-color: var(--border) !important;
    border-radius: 6px !important;
    font-family: 'DM Mono', monospace !important;
}

/* ══════════════════════════════════════════════════════════════
   CUSTOM COMPONENTS
   ══════════════════════════════════════════════════════════════ */

/* ── Page header ────────────────────────────────────────────── */
.page-header {
    background: var(--navy);
    padding: 1.375rem 1.75rem 1.25rem;
    border-radius: 14px;
    margin-bottom: 1.5rem;
    display: flex;
    justify-content: space-between;
    align-items: center;
    position: relative;
    overflow: hidden;
    box-shadow: 0 4px 20px rgba(30,58,95,0.18);
}
.page-header::before {
    content: '';
    position: absolute;
    bottom: 0; left: 0; right: 0;
    height: 3px;
    background: linear-gradient(90deg, var(--accent) 0%, rgba(245,158,11,0.15) 65%, transparent 100%);
}
.page-header::after {
    content: '';
    position: absolute;
    top: -60px; right: -60px;
    width: 200px; height: 200px;
    border-radius: 50%;
    background: radial-gradient(circle, rgba(245,158,11,0.08) 0%, transparent 70%);
    pointer-events: none;
}
.page-header-left { display: flex; align-items: center; gap: 18px; }
.page-header-logo {
    font-family: 'Syne', sans-serif;
    font-size: 1.5rem;
    font-weight: 800;
    color: #fff !important;
    letter-spacing: -0.04em;
    line-height: 1;
}
.page-header-logo span { color: var(--accent); }
.page-header-divider { width: 1px; height: 26px; background: rgba(255,255,255,0.15); }
.page-header-title {
    font-family: 'DM Mono', monospace;
    font-size: 0.68rem;
    color: rgba(255,255,255,0.5);
    letter-spacing: 0.07em;
    text-transform: uppercase;
    line-height: 1.6;
}
.header-badge {
    font-family: 'DM Mono', monospace;
    font-size: 0.66rem;
    font-weight: 500;
    border-radius: 4px;
    padding: 4px 10px;
    letter-spacing: 0.06em;
    text-transform: uppercase;
}
.header-badge.live   { background: rgba(5,150,105,0.2); border: 1px solid rgba(5,150,105,0.4); color: #34D399; }
.header-badge.offline{ background: rgba(245,158,11,0.15); border: 1px solid rgba(245,158,11,0.35); color: #FCD34D; }

/* ── Step tracker ───────────────────────────────────────────── */
.step-tracker {
    display: flex;
    align-items: flex-start;
    background: var(--surface);
    border: 1px solid var(--border);
    border-radius: 10px;
    padding: 1.125rem 2rem 0.875rem;
    margin-bottom: 1.25rem;
    box-shadow: 0 1px 4px rgba(0,0,0,0.05);
}
.step-item { display: flex; flex-direction: column; align-items: center; gap: 6px; }
.step-circle {
    width: 30px; height: 30px; flex-shrink: 0;
    border-radius: 50%;
    display: flex; align-items: center; justify-content: center;
    font-family: 'DM Mono', monospace;
    font-size: 0.68rem; font-weight: 600;
    transition: all 0.2s;
}
.step-circle-done    { background: var(--success); color: #fff; }
.step-circle-active  { background: var(--accent); color: #fff; box-shadow: 0 0 0 4px var(--accent-glow); }
.step-circle-pending { background: var(--surface-3); color: var(--text-dim); border: 1px solid var(--border); }
.step-label { font-family: 'DM Mono', monospace; font-size: 0.6rem; color: var(--text-muted); text-align: center; white-space: nowrap; letter-spacing: 0.05em; }
.step-label-done   { color: var(--success) !important; }
.step-label-active { color: #D97706 !important; font-weight: 600 !important; }
/* margin-top: 14px = (30px circle / 2) - (2px line / 2) → aligns connector with circle center */
.step-conn { height: 2px; flex: 1; margin: 14px 8px 0; min-width: 24px; border-radius: 1px; }
.step-conn-done    { background: var(--success); }
.step-conn-pending { background: var(--border); }

/* ── Section title ──────────────────────────────────────────── */
.section-title {
    font-family: 'DM Mono', monospace;
    font-size: 0.62rem;
    font-weight: 500;
    color: var(--text-muted);
    text-transform: uppercase;
    letter-spacing: 0.14em;
    border-bottom: 1px solid var(--border);
    padding-bottom: 7px;
    margin: 0 0 1rem;
}
.section-title span { color: var(--accent); margin-right: 7px; font-weight: 700; }

/* ── Info card ──────────────────────────────────────────────── */
.info-card {
    background: var(--surface-2);
    border: 1px solid var(--border);
    border-left: 3px solid var(--accent);
    border-radius: 6px;
    padding: 0.75rem 1rem;
    margin-bottom: 0.75rem;
    font-family: 'DM Mono', monospace;
    font-size: 0.78rem;
    color: var(--text);
}
.info-card .label {
    color: var(--text-muted);
    font-size: 0.6rem;
    font-weight: 500;
    text-transform: uppercase;
    letter-spacing: 0.12em;
    margin-bottom: 4px;
}

/* ── Approval box ───────────────────────────────────────────── */
.approval-box {
    padding: 0.625rem 1rem;
    border-radius: 6px;
    text-align: center;
    font-family: 'DM Mono', monospace;
    font-size: 0.72rem;
    font-weight: 500;
    letter-spacing: 0.06em;
    text-transform: uppercase;
    margin: 0.5rem 0;
}
.approval-locked { background: var(--danger-dim); border: 1px solid rgba(220,38,38,0.2); color: var(--danger); }
.approval-ready  { background: var(--success-dim); border: 1px solid rgba(5,150,105,0.25); color: var(--success); }

/* ── Sidebar branding ───────────────────────────────────────── */
.sidebar-logo {
    font-family: 'Syne', sans-serif;
    font-size: 1.35rem;
    font-weight: 800;
    letter-spacing: -0.03em;
    padding: 0.25rem 0 1.25rem;
    border-bottom: 1px solid var(--border);
    margin-bottom: 1.25rem;
    color: var(--navy) !important;
}
.sidebar-logo span { color: var(--accent) !important; }
.sidebar-divider { border-top: 1px solid var(--border); margin: 1rem 0; }
.sidebar-section-label {
    font-family: 'DM Mono', monospace !important;
    font-size: 0.58rem !important;
    font-weight: 500 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.14em !important;
    color: var(--text-muted) !important;
    margin-bottom: 0.4rem;
    display: block;
}

/* ── Top candidate card ─────────────────────────────────────── */
.top-candidate-card {
    background: linear-gradient(135deg, #FFFBEB 0%, #F0FDF4 100%);
    border: 1px solid #FDE68A;
    border-left: 3px solid var(--accent);
    border-radius: 8px;
    padding: 0.875rem 1.125rem;
    margin-top: 0.75rem;
    font-family: 'DM Mono', monospace;
}
.top-candidate-card .tc-label {
    font-size: 0.6rem; font-weight: 600; color: #D97706;
    text-transform: uppercase; letter-spacing: 0.14em; margin-bottom: 6px;
}
.top-candidate-card .tc-name {
    font-family: 'Syne', sans-serif;
    font-size: 0.95rem; font-weight: 700; color: var(--text-bright);
}
.top-candidate-card .tc-meta { font-size: 0.72rem; color: var(--text-muted); margin-top: 5px; line-height: 1.7; }
</style>
"""


# ── Environment ────────────────────────────────────────────────────────────────

def load_env_file(path: Path = ENV_FILE) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and value and key not in os.environ:
            os.environ[key] = value


load_env_file()


# ── Data helpers ───────────────────────────────────────────────────────────────

@st.cache_data(ttl=30)
def load_inventory_from_db() -> pd.DataFrame:
    query = """
        SELECT
            m.material_id,
            m.material_name,
            m.category,
            m.brand,
            m.mpn,
            m.technical_specification,
            i.current_stock,
            i.safety_stock,
            i.plant,
            i.last_updated
        FROM inventory i
        JOIN material_master m ON i.material_id = m.material_id
        ORDER BY m.material_id
    """
    with get_connection() as conn:
        return pd.read_sql_query(query, conn)


def initialize_session() -> None:
    defaults = {
        "step": "idle",
        "approval_ready": False,
        "candidate_results": None,
        "evaluation_report": None,
        "selected_candidate_id": None,
        "selected_material_id": None,
        "po_created": False,
        "po_row": None,
        "last_error": None,
        "messages": [
            {
                "role": "assistant",
                "content": "안녕하세요. 실제 재고 DB 기준으로 부족 자재를 확인하고 웹 리서치를 실행할 수 있습니다.",
            }
        ],
    }
    for key, value in defaults.items():
        st.session_state.setdefault(key, value)


def reset_demo_state() -> None:
    for key in [
        "step",
        "approval_ready",
        "candidate_results",
        "evaluation_report",
        "selected_candidate_id",
        "po_created",
        "po_row",
        "last_error",
    ]:
        st.session_state.pop(key, None)
    st.session_state.messages = [
        {"role": "assistant", "content": "상태를 초기화했습니다. 실제 재고 DB를 다시 읽습니다."}
    ]
    initialize_session()


def clear_pipeline_state_for_material_change(selected_material_id: str) -> None:
    previous_material_id = st.session_state.get("active_pipeline_material_id")
    if previous_material_id in (None, selected_material_id):
        st.session_state.active_pipeline_material_id = selected_material_id
        return

    for key in [
        "candidate_results",
        "evaluation_report",
        "approval_ready",
        "selected_candidate_id",
        "po_row",
        "last_error",
    ]:
        st.session_state.pop(key, None)
    st.session_state.step = "idle"
    st.session_state.active_pipeline_material_id = selected_material_id
    initialize_session()


def shortage_rows(inventory: pd.DataFrame) -> pd.DataFrame:
    shortage = inventory[inventory["current_stock"] < inventory["safety_stock"]].copy()
    shortage["shortage_qty"] = shortage["safety_stock"] - shortage["current_stock"]
    shortage["risk_ratio"] = shortage["current_stock"] / shortage["safety_stock"]
    return shortage.sort_values(["risk_ratio", "shortage_qty"], ascending=[True, False])


@st.cache_data(ttl=30)
def shortage_events_by_material() -> dict[str, dict]:
    events = {}
    for event in scan_inventory():
        events[event.material_id] = json.loads(event.model_dump_json())
    return events


def default_selected_material(shortages: pd.DataFrame) -> str | None:
    if shortages.empty:
        return None
    current = st.session_state.get("selected_material_id")
    if current in set(shortages["material_id"]):
        return current
    return str(shortages.iloc[0]["material_id"])


def write_shortage_event(event: dict) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / f"shortage_event_{event['material_id']}.json"
    path.write_text(json.dumps(event, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def candidate_output_path(material_id: str) -> Path:
    return OUTPUT_DIR / f"candidate_results_{material_id}.json"


def evaluation_output_path(material_id: str) -> Path:
    return OUTPUT_DIR / f"evaluation_report_batch_{material_id}.json"


def run_actual_pipeline(event: dict) -> tuple[dict, dict]:
    event_path = write_shortage_event(event)
    candidate_path = candidate_output_path(event["material_id"])

    has_serpapi = bool(os.environ.get("SERPAPI_API_KEY"))
    has_openai = bool(os.environ.get("OPENAI_API_KEY") or os.environ.get("GPT_API_KEY"))
    provider = "serpapi" if has_serpapi else "llm_plan"

    candidate_results = run_web_research(
        input_path=event_path,
        output_path=candidate_path,
        search_mode="live",
        query_mode="llm" if has_openai else "deterministic",
        extraction_mode="llm" if has_openai else "none",
        search_provider=provider,
        max_results_per_query=3,
        max_candidates=6,
    )

    evaluation_report = evaluate_candidates_from_phase2_results(candidate_path, mode="urgent")
    write_evaluation_report(evaluation_report, evaluation_output_path(event["material_id"]))
    return candidate_results, evaluation_report


def build_candidate_table(candidate_results: dict | None, evaluation_report: dict | None) -> pd.DataFrame:
    if evaluation_report:
        rows = []
        for item in evaluation_report.get("items", []):
            candidate = item.get("candidate_material", {})
            scores = item.get("scores", {})
            decision = item.get("decision_context", {})
            trust_notes = item.get("source_trust_notes", {})
            vendor_notes = item.get("vendor_trust_notes", {})
            rows.append(
                {
                    "candidate_id": candidate.get("candidate_id"),
                    "vendor_name": candidate.get("vendor_name"),
                    "decision": decision.get("decision"),
                    "risk_level": decision.get("risk_level"),
                    "price_krw": candidate.get("price_krw"),
                    "lead_time_days": _display_bool_when_missing(candidate.get("lead_time_days")),
                    "source_type": candidate.get("source_type"),
                    "compatibility_score": scores.get("compatibility_score"),
                    "vendor_trust_score": scores.get("vendor_trust_score"),
                    "source_trust_score": scores.get("source_trust_score"),
                    "final_score": scores.get("final_score"),
                    "risk_note": " / ".join(vendor_notes.get("risk_factors", []) + trust_notes.get("risk_factors", []))
                    or decision.get("recommendation_reason"),
                    "source_url": candidate.get("source_url"),
                }
            )
        return pd.DataFrame(rows)

    rows = []
    for candidate in (candidate_results or {}).get("candidates", []):
        rows.append(
            {
                "candidate_id": candidate.get("candidate_id"),
                "vendor_name": candidate.get("vendor_name"),
                "decision": "not_evaluated",
                "risk_level": "Review",
                "price_krw": candidate.get("price_krw"),
                "lead_time_days": _display_bool_when_missing(candidate.get("lead_time_days")),
                "source_type": candidate.get("source_type"),
                "compatibility_score": None,
                "vendor_trust_score": None,
                "source_trust_score": None,
                "final_score": None,
                "risk_note": candidate.get("spec_evidence"),
                "source_url": candidate.get("source_url"),
            }
        )
    return pd.DataFrame(rows)


def _display_bool_when_missing(value):
    return False if value is None else value


def get_top_evaluation_item(evaluation_report: dict | None) -> dict | None:
    if not evaluation_report:
        return None
    items = evaluation_report.get("items", [])
    top_id = evaluation_report.get("top_candidate_id")
    return next(
        (item for item in items if item.get("candidate_material", {}).get("candidate_id") == top_id),
        items[0] if items else None,
    )


def get_evaluation_item_by_candidate_id(evaluation_report: dict | None, candidate_id: str | None) -> dict | None:
    if not evaluation_report or not candidate_id:
        return None
    return next(
        (
            item for item in evaluation_report.get("items", [])
            if str(item.get("candidate_material", {}).get("candidate_id")) == str(candidate_id)
        ),
        None,
    )


def _positive_int(value) -> int | None:
    try:
        number = int(float(value))
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def is_candidate_po_selectable(item: dict | None) -> bool:
    if not item:
        return False
    decision = str(item.get("decision_context", {}).get("decision", "")).lower()
    candidate = item.get("candidate_material", {})
    return (
        decision in {"recommend", "conditional_approve", "review_required"}
        and _positive_int(candidate.get("price_krw")) is not None
    )


def report_with_selected_candidate(evaluation_report: dict, candidate_id: str) -> dict:
    item = get_evaluation_item_by_candidate_id(evaluation_report, candidate_id)
    if not item:
        raise RuntimeError(f"Selected candidate not found: {candidate_id}")
    if not is_candidate_po_selectable(item):
        decision = item.get("decision_context", {}).get("decision", "unknown")
        price = item.get("candidate_material", {}).get("price_krw")
        if _positive_int(price) is None:
            raise RuntimeError("선택한 후보에 확정 단가가 없어 PO를 생성할 수 없습니다.")
        raise RuntimeError(f"Selected candidate cannot be converted to PO: {decision}")
    selected_report = dict(evaluation_report)
    selected_report["top_candidate_id"] = candidate_id
    selected_report["next_action"] = "approval_pending"
    return selected_report


def run_phase4_approval(event: dict, evaluation_report: dict, selected_candidate_id: str) -> dict:
    graph = build_graph()
    state = make_initial_state(f"streamlit-{event['material_id']}-{datetime.now().strftime('%H%M%S')}")
    state["shortage_event"] = {
        "material_id": event.get("material_id"),
        "description": event.get("material_name"),
        "shortage_qty": event.get("shortage_qty", 0),
    }
    state["evaluation_report_batch"] = report_with_selected_candidate(evaluation_report, selected_candidate_id)

    config = {"configurable": {"thread_id": state["workflow_id"]}}
    for _ in graph.stream(state, config=config):
        pass

    current_state = graph.get_state(config).values
    if current_state.get("status") != "APPROVAL_PENDING":
        raise RuntimeError(f"Phase 4 approval is not available: {current_state.get('status')}")

    current_state["approval"] = {
        "required": True,
        "approved": True,
        "approver": "streamlit_user",
        "approved_at": datetime.now().astimezone().isoformat(timespec="seconds"),
    }
    graph.update_state(config, current_state)
    for _ in graph.stream(None, config=config):
        pass
    return graph.get_state(config).values


def write_po_from_phase4_state(final_state: dict, evaluation_report: dict, selected_candidate_id: str) -> dict:
    po_draft = final_state.get("po_draft")
    if not po_draft:
        raise RuntimeError("Phase 4 did not produce po_draft.")

    item = get_evaluation_item_by_candidate_id(evaluation_report, selected_candidate_id) or {}
    candidate = item.get("candidate_material", {})
    scores = item.get("scores", {})
    decision = item.get("decision_context", {})
    po_row = {
        "PO_DRAFT_NO": f"PO-DRAFT-{datetime.now().strftime('%Y%m%d%H%M%S')}",
        "CREATED_AT": datetime.now().isoformat(timespec="seconds"),
        "WORKFLOW_STATUS": final_state.get("status"),
        "SOURCE": "phase4_po_draft",
        "VENDOR_NAME": po_draft.get("vendor_name"),
        "MATERIAL_ID": po_draft.get("material_code"),
        "CANDIDATE_ID": candidate.get("candidate_id"),
        "UNIT_PRICE_KRW": po_draft.get("unit_price"),
        "ORDER_QTY": po_draft.get("quantity"),
        "TOTAL_AMOUNT_KRW": po_draft.get("total_price"),
        "LEAD_TIME_DAYS": candidate.get("lead_time_days"),
        "SOURCE_URL": candidate.get("source_url"),
        "FINAL_SCORE": scores.get("final_score"),
        "RISK_NOTE": decision.get("recommendation_reason"),
    }
    pd.DataFrame([po_row]).to_csv(PO_FILE, index=False)
    st.session_state.po_created = True
    st.session_state.po_row = po_row
    return po_row


# ── Assistant logic ────────────────────────────────────────────────────────────

def assistant_reply(prompt: str, event: dict) -> str:
    text = prompt.strip().lower()

    if any(word in text for word in ["재고", "위험", "부족", "확인"]):
        st.session_state.step = "inventory_checked"
        st.session_state.approval_ready = False
        return (
            f"실제 DB 기준 **{event['material_id']} ({event['material_name']})** 재고는 "
            f"{event['current_stock']}개이고 최소 보유 수량은 {event['safety_stock']}개입니다. "
            f"부족 수량은 **{event['shortage_qty']}개**입니다."
        )

    if any(word in text for word in ["대체", "찾아", "검색", "후보", "업체"]):
        try:
            with st.spinner("웹 리서치와 평가를 실행하는 중입니다..."):
                candidate_results, evaluation_report = run_actual_pipeline(event)
        except Exception as exc:
            st.session_state.last_error = str(exc)
            st.session_state.approval_ready = False
            return f"파이프라인 실행이 실패했습니다.\n\n`{exc}`"

        st.session_state.candidate_results = candidate_results
        st.session_state.evaluation_report = evaluation_report
        st.session_state.active_pipeline_material_id = event["material_id"]
        st.session_state.step = "candidates_loaded"
        st.session_state.selected_candidate_id = None
        st.session_state.po_created = False
        st.session_state.po_row = None

        table = build_candidate_table(candidate_results, evaluation_report)
        action = evaluation_report.get("next_action")
        st.session_state.approval_ready = action == "approval_pending"
        top_item = get_top_evaluation_item(evaluation_report)
        top_candidate = top_item.get("candidate_material", {}) if top_item else {}

        if action == "approval_pending":
            return (
                f"후보 {len(table)}개를 수집하고 Phase 3/4 승인 대기 조건을 확인했습니다. "
                f"승인 대상은 **{top_candidate.get('vendor_name')} / {top_candidate.get('candidate_id')}**입니다."
            )
        if action == "manual_review":
            st.session_state.approval_ready = True
            return "후보를 찾았습니다. 후보 표 아래에서 최종 후보를 선택하면 PO를 생성할 수 있습니다."
        return "PO로 전환할 수 있는 후보를 찾지 못했습니다. reject 후보만 있는지 평가 결과를 확인해 주세요."

    if any(word in text for word in ["반려", "거절", "보류", "중단"]):
        st.session_state.step = "rejected"
        st.session_state.approval_ready = False
        return "반려로 기록했습니다. PO 초안은 생성하지 않았습니다."

    if any(word in text for word in ["승인", "발주", "진행", "ok", "yes", "응"]):
        if not st.session_state.approval_ready or not st.session_state.evaluation_report:
            return "아직 후보 평가가 완료되지 않았습니다. 먼저 `대체품 찾아줘`로 후보 검색과 평가를 실행해주세요."
        try:
            selected_candidate_id = st.session_state.get("selected_candidate_id")
            if not selected_candidate_id:
                return "최종 후보를 먼저 선택해주세요. 후보 표 아래 선택 박스에서 PO를 생성할 후보를 고를 수 있습니다."
            final_state = run_phase4_approval(event, st.session_state.evaluation_report, selected_candidate_id)
            po_row = write_po_from_phase4_state(final_state, st.session_state.evaluation_report, selected_candidate_id)
        except Exception as exc:
            return f"PO 초안 생성에 실패했습니다.\n\n`{exc}`"
        st.session_state.step = "po_created"
        st.session_state.approval_ready = False
        return (
            f"선택한 후보 기준으로 `{PO_FILE.name}`를 생성했습니다.\n\n"
            f"{po_row['PO_DRAFT_NO']} / {po_row['VENDOR_NAME']} / "
            f"{po_row['ORDER_QTY']}개 / {po_row['TOTAL_AMOUNT_KRW']:,}원"
        )

    return "재고 확인, 대체품 검색, 승인, 반려 중 하나로 진행할 수 있습니다."


# ── UI rendering ───────────────────────────────────────────────────────────────

def _step_index(step: str, approval_ready: bool) -> int:
    if step == "idle":
        return 0
    if step == "inventory_checked":
        return 1
    if step in ("candidates_loaded", "rejected"):
        return 3 if approval_ready else 2
    if step == "po_created":
        return 4
    return 0


def render_page_header(has_serpapi: bool) -> None:
    mode_cls = "live" if has_serpapi else "offline"
    mode_lbl = "● LIVE · SerpAPI" if has_serpapi else "● OFFLINE · LLM Plan"
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    st.markdown(
        f"""
        <div class="page-header">
            <div class="page-header-left">
                <div class="page-header-logo">Buy<span>Bee</span></div>
                <div class="page-header-divider"></div>
                <div class="page-header-title">
                    Procurement Intelligence Platform<br>
                    부족 자재 탐지 · 대체 후보 비교 · 발주 문서 생성
                </div>
            </div>
            <div style="display:flex;align-items:center;gap:12px;">
                <div class="header-badge {mode_cls}">{mode_lbl}</div>
                <div style="font-family:'DM Mono',monospace;font-size:0.64rem;color:rgba(255,255,255,0.32);letter-spacing:0.04em;">{now}</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_step_tracker(step: str, approval_ready: bool) -> None:
    labels = ["재고 분석", "웹 리서치", "후보 평가", "승인 대기", "PO 완료"]
    current = _step_index(step, approval_ready)

    items_html = ""
    for i, label in enumerate(labels):
        if i < current:
            circle_cls, label_cls, icon = "step-circle-done", "step-label-done", "✓"
        elif i == current:
            circle_cls, label_cls, icon = "step-circle-active", "step-label-active", str(i + 1)
        else:
            circle_cls, label_cls, icon = "step-circle-pending", "", str(i + 1)

        items_html += f"""
        <div class="step-item">
            <div class="step-circle {circle_cls}">{icon}</div>
            <div class="step-label {label_cls}">{label}</div>
        </div>
        """
        if i < len(labels) - 1:
            conn_cls = "step-conn-done" if i < current else "step-conn-pending"
            items_html += f'<div class="step-conn {conn_cls}"></div>'

    st.markdown(f'<div class="step-tracker">{items_html}</div>', unsafe_allow_html=True)


def render_sidebar(shortages: pd.DataFrame, selected_default: str | None) -> str:
    has_serpapi = bool(os.environ.get("SERPAPI_API_KEY"))
    has_openai = bool(os.environ.get("OPENAI_API_KEY") or os.environ.get("GPT_API_KEY"))

    with st.container(border=True):
        st.markdown('<div class="sidebar-logo">Buy<span>Bee</span> 🐝</div>', unsafe_allow_html=True)

        st.markdown('<span class="sidebar-section-label">부족 자재 선택</span>', unsafe_allow_html=True)
        selected_material = st.selectbox(
            "부족 자재",
            list(shortages["material_id"]),
            index=list(shortages["material_id"]).index(selected_default),
            label_visibility="collapsed",
        )

        shortage_row = shortages[shortages["material_id"] == selected_material].iloc[0]
        risk_pct = int(shortage_row["risk_ratio"] * 100)
        risk_color = "#DC2626" if risk_pct < 30 else "#D97706" if risk_pct < 70 else "#059669"
        st.markdown(
            f"""
            <div style="background:rgba(30,58,95,0.035);border:1px solid #E4E8F0;
                        border-left:3px solid {risk_color};border-radius:6px;
                        padding:0.75rem 1rem;margin:0.5rem 0 1rem;
                        font-family:'DM Mono',monospace;font-size:0.75rem;">
                <div style="color:#6B7280;font-size:0.58rem;font-weight:500;
                            text-transform:uppercase;letter-spacing:0.14em;margin-bottom:8px;">재고 현황</div>
                <div style="display:flex;justify-content:space-between;margin-bottom:5px;color:#6B7280;">
                    <span>현재 재고</span>
                    <span style="color:#111827;font-weight:500;">{int(shortage_row['current_stock'])}개</span>
                </div>
                <div style="display:flex;justify-content:space-between;margin-bottom:10px;color:#6B7280;">
                    <span>안전 재고</span>
                    <span style="color:#111827;font-weight:500;">{int(shortage_row['safety_stock'])}개</span>
                </div>
                <div style="background:#E4E8F0;border-radius:2px;height:3px;overflow:hidden;">
                    <div style="background:{risk_color};width:{min(risk_pct,100)}%;height:100%;"></div>
                </div>
                <div style="color:{risk_color};font-size:0.62rem;margin-top:5px;text-align:right;letter-spacing:0.06em;">
                    재고율 {risk_pct}%
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.markdown('<div class="sidebar-divider"></div>', unsafe_allow_html=True)
        st.markdown('<span class="sidebar-section-label">API 연동 상태</span>', unsafe_allow_html=True)

        s_color = "#34D399" if has_serpapi else "#F87171"
        o_color = "#34D399" if has_openai else "#F87171"
        st.markdown(
            f"""
            <div style="font-family:'DM Mono',monospace;font-size:0.73rem;line-height:2.2;color:#6B7280;">
                <span style="color:{s_color};">●</span>&nbsp; SerpAPI &nbsp;
                <span style="color:{s_color};font-size:0.62rem;">{"CONNECTED" if has_serpapi else "OFFLINE"}</span><br>
                <span style="color:{o_color};">●</span>&nbsp; OpenAI &nbsp;&nbsp;
                <span style="color:{o_color};font-size:0.62rem;">{"CONNECTED" if has_openai else "OFFLINE"}</span>
            </div>
            """,
            unsafe_allow_html=True,
        )
        if not has_serpapi:
            st.info("SERPAPI_API_KEY 없음 — 오프라인 계획 모드")

        st.markdown('<div class="sidebar-divider"></div>', unsafe_allow_html=True)
        st.markdown('<span class="sidebar-section-label">예산 한도</span>', unsafe_allow_html=True)
        st.markdown(
            f'<div style="font-family:\'Syne\',sans-serif;font-size:1.1rem;font-weight:700;color:#F59E0B;">{BUDGET_LIMIT_KRW:,}원</div>',
            unsafe_allow_html=True,
        )

    return selected_material


def render_top_metrics(event: dict) -> None:
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("자재 ID", event["material_id"])
    c2.metric("카테고리", event["category"])
    c3.metric("부족 수량", f"{event['shortage_qty']:,} 개")

    step = st.session_state.step
    approval_ready = st.session_state.approval_ready
    if step == "po_created":
        status_val, status_delta = "PO 완료", None
    elif approval_ready:
        status_val, status_delta = "승인 대기", "준비됨"
    elif step == "candidates_loaded":
        status_val, status_delta = "평가 완료", None
    elif step == "inventory_checked":
        status_val, status_delta = "재고 확인됨", None
    else:
        status_val, status_delta = "대기 중", None

    c4.metric("워크플로 상태", status_val, status_delta)


def render_inventory(inventory: pd.DataFrame, shortages: pd.DataFrame) -> None:
    st.markdown('<div class="section-title"><span>01</span>실시간 재고 현황</div>', unsafe_allow_html=True)

    chart = (
        inventory[["material_name", "current_stock", "safety_stock"]]
        .rename(columns={
            "current_stock": "현재 재고",
            "safety_stock": "최소 보유 수량",
        })
        .set_index("material_name")
    )
    st.bar_chart(chart, color=["#DC2626", "#2563EB"], use_container_width=True)

    display_cols = [
        "material_id", "material_name", "category", "brand", "mpn",
        "current_stock", "safety_stock", "shortage_qty", "plant",
    ]
    st.dataframe(
        shortages[display_cols].rename(columns={
            "material_id": "자재 ID",
            "material_name": "자재명",
            "category": "카테고리",
            "brand": "브랜드",
            "mpn": "MPN",
            "current_stock": "현재 재고",
            "safety_stock": "최소 보유 수량",
            "shortage_qty": "부족 수량",
            "plant": "플랜트",
        }),
        use_container_width=True,
        hide_index=True,
    )


def render_pipeline_results(event: dict) -> None:
    st.markdown('<div class="section-title"><span>02</span>웹 리서치 · 후보 평가 결과</div>', unsafe_allow_html=True)

    candidate_results = st.session_state.candidate_results
    evaluation_report = st.session_state.evaluation_report

    if st.session_state.last_error:
        st.error(st.session_state.last_error)

    if not candidate_results:
        st.info("아직 웹 검색을 실행하지 않았습니다. 우측 어시스턴트에서 **대체품 찾아줘**를 입력하세요.")
        return

    if candidate_results.get("material_id") != event.get("material_id"):
        st.warning("선택한 부족 자재와 검색 결과가 달라 결과를 초기화했습니다. 다시 검색을 실행하세요.")
        st.session_state.candidate_results = None
        st.session_state.evaluation_report = None
        st.session_state.approval_ready = False
        st.session_state.step = "idle"
        st.rerun()

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("검색 ID", candidate_results.get("search_id", "-"))
    c2.metric("후보 수", len(candidate_results.get("candidates", [])))
    next_action = evaluation_report.get("next_action", "-") if evaluation_report else "-"
    c3.metric("다음 액션", next_action)
    top_item = get_top_evaluation_item(evaluation_report)
    top_score = top_item.get("scores", {}).get("final_score", "-") if top_item else "-"
    c4.metric("최고 점수", top_score)

    st.markdown(
        f"""
        <div class="info-card">
            <div class="label">부족 자재 스펙</div>
            <strong>{event.get("material_id")} · {event.get("material_name")}</strong>
            <div style="color:var(--muted);font-size:0.8rem;margin-top:4px;">
                {event.get("technical_specification", "—")}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    table = build_candidate_table(candidate_results, evaluation_report)
    if table.empty:
        st.warning("검증된 후보가 없습니다. 검색 키워드 또는 API 설정을 확인하세요.")
        return

    st.dataframe(
        table[[
            "candidate_id", "vendor_name", "decision", "risk_level",
            "price_krw", "lead_time_days", "source_type",
            "compatibility_score", "vendor_trust_score", "source_trust_score", "final_score",
            "risk_note", "source_url",
        ]],
        use_container_width=True,
        hide_index=True,
    )

    evaluated_items = (evaluation_report or {}).get("items", [])
    po_selectable_items = [
        item for item in (evaluation_report or {}).get("items", [])
        if is_candidate_po_selectable(item)
    ]
    price_missing_items = [
        item for item in evaluated_items
        if str(item.get("decision_context", {}).get("decision", "")).lower()
        in {"recommend", "conditional_approve", "review_required"}
        and _positive_int(item.get("candidate_material", {}).get("price_krw")) is None
    ]

    if top_item:
        candidate = top_item.get("candidate_material", {})
        decision = top_item.get("decision_context", {})
        scores = top_item.get("scores", {})
        url = candidate.get("source_url", "")
        url_html = f'<a href="{url}" target="_blank" style="color:var(--accent);font-size:0.78rem;">{url}</a>' if url else "—"
        st.markdown(
            f"""
            <div class="top-candidate-card">
                <div class="tc-label">평가 1순위 후보</div>
                <div class="tc-name">{candidate.get("vendor_name", "—")} &nbsp;·&nbsp; {candidate.get("candidate_id", "—")}</div>
                <div class="tc-meta">
                    결정: <strong>{decision.get("decision", "—")}</strong> &nbsp;|&nbsp;
                    리스크: <strong>{decision.get("risk_level", "—")}</strong> &nbsp;|&nbsp;
                    최종 점수: <strong>{scores.get("final_score", "—")}</strong>
                </div>
                <div style="margin-top:6px;">{url_html}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    if po_selectable_items:
        options = [item.get("candidate_material", {}).get("candidate_id") for item in po_selectable_items]
        labels = {}
        for item in po_selectable_items:
            candidate = item.get("candidate_material", {})
            scores = item.get("scores", {})
            decision = item.get("decision_context", {})
            cid = candidate.get("candidate_id")
            labels[cid] = (
                f"{cid} | {candidate.get('vendor_name', '-')} | "
                f"{decision.get('decision', '-')} | score {scores.get('final_score', '-')}"
            )
        current = st.session_state.get("selected_candidate_id")
        index = options.index(current) if current in options else 0
        selected = st.selectbox(
            "최종 후보 선택",
            options=options,
            index=index,
            format_func=lambda value: labels.get(value, str(value)),
            key="selected_candidate_id",
        )
        if st.button("선택한 후보로 PO 생성", use_container_width=True):
            try:
                final_state = run_phase4_approval(event, evaluation_report, selected)
                po_row = write_po_from_phase4_state(final_state, evaluation_report, selected)
            except Exception as exc:
                st.error(f"PO 생성에 실패했습니다: {exc}")
                return
            st.session_state.step = "po_created"
            st.session_state.approval_ready = False
            st.success(
                f"PO 생성 완료: {po_row['PO_DRAFT_NO']} / {po_row['VENDOR_NAME']} / "
                f"{po_row['ORDER_QTY']}개 / {int(po_row['TOTAL_AMOUNT_KRW']):,}원"
            )
            st.rerun()
    elif evaluation_report:
        if price_missing_items:
            st.warning("PO로 전환 가능한 평가 후보는 있지만 확정 단가가 없습니다. 단가가 확인된 후보를 선택하거나 가격 추출 후 다시 실행해 주세요.")
        else:
            st.info("PO로 전환할 수 있는 후보가 없습니다. reject 후보만 있는지 평가 결과를 확인해 주세요.")

    if False and top_item:
        candidate = top_item.get("candidate_material", {})
        decision = top_item.get("decision_context", {})
        scores = top_item.get("scores", {})
        url = candidate.get("source_url", "")
        url_html = f'<a href="{url}" target="_blank" style="color:var(--accent);font-size:0.78rem;">{url}</a>' if url else "—"
        st.markdown(
            f"""
            <div class="top-candidate-card">
                <div class="tc-label">최우선 후보</div>
                <div class="tc-name">{candidate.get("vendor_name", "—")} &nbsp;·&nbsp; {candidate.get("candidate_id", "—")}</div>
                <div class="tc-meta">
                    결정: <strong>{decision.get("decision", "—")}</strong> &nbsp;|&nbsp;
                    리스크: <strong>{decision.get("risk_level", "—")}</strong> &nbsp;|&nbsp;
                    최종 점수: <strong>{scores.get("final_score", "—")}</strong>
                </div>
                <div style="margin-top:6px;">{url_html}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )


def build_excel_report(event: dict) -> bytes:
    """PO 초안 + 후보 평가 + 부족 자재 정보를 담은 스타일된 Excel 보고서를 생성한다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    NAVY   = "1E3A5F"
    AMBER  = "F59E0B"
    LIGHT  = "F3F5F9"
    WHITE  = "FFFFFF"
    GREEN  = "059669"
    RED    = "DC2626"
    BORDER_COLOR = "D0D7E6"

    thin = Side(style="thin", color=BORDER_COLOR)
    full_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr_font(color=WHITE, size=10, bold=True):
        return Font(name="Calibri", bold=bold, color=color, size=size)

    def hdr_fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def center(wrap=False):
        return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

    def left(wrap=False):
        return Alignment(horizontal="left", vertical="center", wrap_text=wrap)

    wb = Workbook()

    # ── Sheet 1: PO 초안 ──────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "PO 초안"

    # 타이틀 영역
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "BuyBee — 발주 초안 보고서"
    ws1["A1"].font = Font(name="Calibri", bold=True, size=14, color=WHITE)
    ws1["A1"].fill = hdr_fill(NAVY)
    ws1["A1"].alignment = center()

    ws1.merge_cells("A2:G2")
    ws1["A2"] = f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws1["A2"].font = Font(name="Calibri", size=9, color="94A3B8")
    ws1["A2"].fill = hdr_fill("162D4A")
    ws1["A2"].alignment = left()

    ws1.row_dimensions[1].height = 28
    ws1.row_dimensions[2].height = 16

    # PO 데이터
    if PO_FILE.exists():
        po_df = pd.read_csv(PO_FILE)
        if not po_df.empty:
            headers = list(po_df.columns)
            ws1.append([])  # 빈 줄
            ws1.append(headers)
            hdr_row = ws1.max_row
            for col_idx, _ in enumerate(headers, 1):
                cell = ws1.cell(row=hdr_row, column=col_idx)
                cell.font = hdr_font(WHITE)
                cell.fill = hdr_fill(NAVY)
                cell.alignment = center()
                cell.border = full_border

            for _, data_row in po_df.iterrows():
                ws1.append(list(data_row))
                r = ws1.max_row
                for col_idx in range(1, len(headers) + 1):
                    cell = ws1.cell(row=r, column=col_idx)
                    cell.fill = hdr_fill(LIGHT)
                    cell.alignment = left(wrap=True)
                    cell.border = full_border
                    cell.font = Font(name="Calibri", size=9)

    # 열 너비 (병합 셀 제외)
    for col in ws1.columns:
        first = next((c for c in col if hasattr(c, "column_letter")), None)
        if first is None:
            continue
        max_len = max((len(str(c.value or "")) for c in col if hasattr(c, "column_letter")), default=8)
        ws1.column_dimensions[first.column_letter].width = min(max_len + 4, 40)

    # ── Sheet 2: 후보 평가 결과 ───────────────────────────────────
    ws2 = wb.create_sheet("후보 평가 결과")

    ws2.merge_cells("A1:L1")
    ws2["A1"] = "후보 평가 결과 — 전체 목록"
    ws2["A1"].font = Font(name="Calibri", bold=True, size=13, color=WHITE)
    ws2["A1"].fill = hdr_fill(NAVY)
    ws2["A1"].alignment = center()
    ws2.row_dimensions[1].height = 26

    evaluation_report = st.session_state.get("evaluation_report")
    candidate_results = st.session_state.get("candidate_results")
    table = build_candidate_table(candidate_results, evaluation_report)

    if not table.empty:
        ws2.append([])
        col_labels = {
            "candidate_id": "후보 ID", "vendor_name": "공급업체",
            "decision": "결정", "risk_level": "리스크",
            "price_krw": "단가(원)", "lead_time_days": "납기(일)",
            "source_type": "출처 유형",
            "compatibility_score": "호환성", "vendor_trust_score": "공급사 신뢰도",
            "source_trust_score": "출처 신뢰도",
            "final_score": "최종점수", "risk_note": "리스크 메모",
            "source_url": "출처 URL",
        }
        headers = [col_labels.get(c, c) for c in table.columns]
        ws2.append(headers)
        hdr_row = ws2.max_row
        for col_idx, _ in enumerate(headers, 1):
            cell = ws2.cell(row=hdr_row, column=col_idx)
            cell.font = hdr_font(WHITE)
            cell.fill = hdr_fill(NAVY)
            cell.alignment = center()
            cell.border = full_border

        top_id = (evaluation_report or {}).get("top_candidate_id")
        for _, data_row in table.iterrows():
            ws2.append(list(data_row))
            r = ws2.max_row
            is_top = str(data_row.get("candidate_id")) == str(top_id)
            fill_color = "FFFBEB" if is_top else WHITE
            for col_idx in range(1, len(headers) + 1):
                cell = ws2.cell(row=r, column=col_idx)
                cell.fill = hdr_fill(fill_color)
                cell.alignment = left(wrap=True)
                cell.border = full_border
                cell.font = Font(name="Calibri", size=9,
                                 bold=is_top, color="92400E" if is_top else "374151")

            # 결정 컬럼 색상
            decision_col = list(table.columns).index("decision") + 1
            risk_col = list(table.columns).index("risk_level") + 1
            decision_val = str(data_row.get("decision", ""))
            risk_val = str(data_row.get("risk_level", ""))
            d_cell = ws2.cell(row=r, column=decision_col)
            r_cell = ws2.cell(row=r, column=risk_col)
            if "approve" in decision_val.lower():
                d_cell.font = Font(name="Calibri", size=9, bold=True, color=GREEN)
            elif "reject" in decision_val.lower():
                d_cell.font = Font(name="Calibri", size=9, bold=True, color=RED)
            if "high" in risk_val.lower():
                r_cell.font = Font(name="Calibri", size=9, bold=True, color=RED)
            elif "low" in risk_val.lower():
                r_cell.font = Font(name="Calibri", size=9, bold=True, color=GREEN)

        for col in ws2.columns:
            first = next((c for c in col if hasattr(c, "column_letter")), None)
            if first is None:
                continue
            max_len = max((len(str(c.value or "")) for c in col if hasattr(c, "column_letter")), default=8)
            ws2.column_dimensions[first.column_letter].width = min(max_len + 4, 50)

    # ── Sheet 3: 부족 자재 정보 ───────────────────────────────────
    ws3 = wb.create_sheet("부족 자재 정보")

    ws3.merge_cells("A1:B1")
    ws3["A1"] = "부족 자재 상세 정보"
    ws3["A1"].font = Font(name="Calibri", bold=True, size=13, color=WHITE)
    ws3["A1"].fill = hdr_fill(NAVY)
    ws3["A1"].alignment = center()
    ws3.row_dimensions[1].height = 26
    ws3.column_dimensions["A"].width = 24
    ws3.column_dimensions["B"].width = 52

    fields = [
        ("자재 ID",         event.get("material_id", "")),
        ("자재명",          event.get("material_name", "")),
        ("카테고리",        event.get("category", "")),
        ("브랜드",          event.get("brand", "")),
        ("MPN",            event.get("mpn", "")),
        ("현재 재고",       event.get("current_stock", "")),
        ("최소 보유 수량",  event.get("safety_stock", "")),
        ("부족 수량",       event.get("shortage_qty", "")),
        ("기술 스펙",       event.get("technical_specification", "")),
        ("플랜트",          event.get("plant", "")),
        ("보고서 생성일",   datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    ws3.append([])
    for label, value in fields:
        ws3.append([label, value])
        r = ws3.max_row
        a_cell = ws3.cell(row=r, column=1)
        b_cell = ws3.cell(row=r, column=2)
        a_cell.font = Font(name="Calibri", bold=True, size=9, color=NAVY)
        a_cell.fill = hdr_fill(LIGHT)
        a_cell.alignment = left()
        a_cell.border = full_border
        b_cell.font = Font(name="Calibri", size=9)
        b_cell.alignment = left(wrap=True)
        b_cell.border = full_border
        ws3.row_dimensions[r].height = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _po_money(value) -> str:
    try:
        return f"KRW {int(float(value)):,}"
    except (TypeError, ValueError):
        return "KRW 0"


def build_po_word_document(event: dict, row: pd.Series | dict) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(f"Purchase Order no: {row.get('PO_DRAFT_NO', '-')}")
    doc.add_paragraph(f"Date of issue: {datetime.now().strftime('%Y-%m-%d')}")
    doc.add_paragraph("")
    doc.add_paragraph(f"Supplier: {row.get('VENDOR_NAME', '-')}")
    doc.add_paragraph("Buyer: BuyBee Procurement")
    doc.add_paragraph("")
    doc.add_paragraph("ITEMS")
    table = doc.add_table(rows=2, cols=8)
    table.style = "Table Grid"
    headers = ["No.", "Description", "Qty", "UM", "Net price", "Net worth", "VAT [%]", "Gross worth"]
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    qty = int(row.get("ORDER_QTY", 0) or 0)
    unit_price = int(row.get("UNIT_PRICE_KRW", 0) or 0)
    net = qty * unit_price
    gross = round(net * 1.1)
    values = [
        "1.",
        f"{event.get('material_name', '-')}\nMaterial ID: {row.get('MATERIAL_ID', '-')}\nCandidate ID: {row.get('CANDIDATE_ID', '-')}",
        str(qty),
        "each",
        _po_money(unit_price),
        _po_money(net),
        "10%",
        _po_money(gross),
    ]
    for idx, value in enumerate(values):
        table.cell(1, idx).text = value
    doc.add_paragraph("")
    doc.add_paragraph("SUMMARY")
    summary = doc.add_table(rows=3, cols=4)
    summary.style = "Table Grid"
    for idx, header in enumerate(["VAT [%]", "Net worth", "VAT", "Gross worth"]):
        summary.cell(0, idx).text = header
    for idx, value in enumerate(["10%", _po_money(net), _po_money(round(net * 0.1)), _po_money(gross)]):
        summary.cell(1, idx).text = value
    for idx, value in enumerate(["Total", _po_money(net), _po_money(round(net * 0.1)), _po_money(gross)]):
        summary.cell(2, idx).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_po_pdf_document(event: dict, row: pd.Series | dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = "Helvetica"
    font_bold_name = "Helvetica-Bold"
    malgun = Path(r"C:\Windows\Fonts\malgun.ttf")
    malgun_bold = Path(r"C:\Windows\Fonts\malgunbd.ttf")
    if malgun.exists():
        font_name = "MalgunGothic"
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_name, str(malgun)))
    if malgun_bold.exists():
        font_bold_name = "MalgunGothic-Bold"
        if font_bold_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_bold_name, str(malgun_bold)))

    qty = int(row.get("ORDER_QTY", 0) or 0)
    unit_price = int(row.get("UNIT_PRICE_KRW", 0) or 0)
    net = qty * unit_price
    vat = round(net * 0.1)
    gross = net + vat
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm)
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontName = font_name
    styles["Heading2"].fontName = font_bold_name
    story = [
        Paragraph(f"<b>Purchase Order no:</b> {row.get('PO_DRAFT_NO', '-')}", body),
        Paragraph(f"Date of issue: {datetime.now().strftime('%Y-%m-%d')}", body),
        Spacer(1, 28 * mm),
        Paragraph(f"<b>Supplier:</b> {row.get('VENDOR_NAME', '-')} &nbsp;&nbsp;&nbsp; <b>Buyer:</b> BuyBee Procurement", body),
        Spacer(1, 10 * mm),
        Paragraph("<b>ITEMS</b>", styles["Heading2"]),
    ]
    item_data = [
        ["No.", "Description", "Qty", "UM", "Net price", "Net worth", "VAT [%]", "Gross worth"],
        [
            "1.",
            Paragraph(f"{event.get('material_name', '-')}<br/>Material ID: {row.get('MATERIAL_ID', '-')}<br/>Candidate ID: {row.get('CANDIDATE_ID', '-')}", body),
            str(qty),
            "each",
            _po_money(unit_price),
            _po_money(net),
            "10%",
            _po_money(gross),
        ],
    ]
    item_table = Table(item_data, colWidths=[10*mm, 58*mm, 12*mm, 13*mm, 22*mm, 24*mm, 16*mm, 25*mm])
    item_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#E6E6E6")),
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTNAME", (0, 0), (-1, 0), font_bold_name),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.extend([item_table, Spacer(1, 10 * mm), Paragraph("<b>SUMMARY</b>", styles["Heading2"])])
    summary = Table([
        ["VAT [%]", "Net worth", "VAT", "Gross worth"],
        ["10%", _po_money(net), _po_money(vat), _po_money(gross)],
        ["Total", _po_money(net), _po_money(vat), _po_money(gross)],
    ], colWidths=[42*mm, 42*mm, 42*mm, 42*mm])
    summary.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#E6E6E6")),
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTNAME", (0, 0), (-1, 0), font_bold_name),
        ("FONTNAME", (0, 2), (-1, 2), font_bold_name),
        ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
    ]))
    story.append(summary)
    doc.build(story)
    return buf.getvalue()


def render_po(event: dict) -> None:
    if not PO_FILE.exists():
        return
    st.markdown('<div class="section-title"><span>03</span>최종 발주 보고서</div>', unsafe_allow_html=True)
    po_df = pd.read_csv(PO_FILE)

    if not po_df.empty:
        row = po_df.iloc[0]
        cols = st.columns(4)
        cols[0].metric("PO 번호", str(row.get("PO_DRAFT_NO", "—"))[-12:])
        cols[1].metric("공급업체", str(row.get("VENDOR_NAME", "—")))
        cols[2].metric("발주 수량", f"{int(row.get('ORDER_QTY', 0)):,}개" if pd.notna(row.get("ORDER_QTY")) else "—")
        total = row.get("TOTAL_AMOUNT_KRW")
        cols[3].metric("총액", f"{int(total):,}원" if pd.notna(total) else "—")

        filename = f"BuyBee_PO_{row.get('MATERIAL_ID', 'report')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        docx_filename = f"BuyBee_PO_{row.get('MATERIAL_ID', 'report')}_{datetime.now().strftime('%Y%m%d')}.docx"
        pdf_filename = f"BuyBee_PO_{row.get('MATERIAL_ID', 'report')}_{datetime.now().strftime('%Y%m%d')}.pdf"
        excel_bytes = build_excel_report(event)
        word_bytes = build_po_word_document(event, row)
        pdf_bytes = build_po_pdf_document(event, row)
        c_excel, c_word, c_pdf = st.columns(3)
        c_excel.download_button(
            label="Excel 보고서 다운로드",
            data=excel_bytes,
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
        c_word.download_button(
            label="PO Word 문서 다운로드",
            data=word_bytes,
            file_name=docx_filename,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )
        c_pdf.download_button(
            label="PO PDF 문서 다운로드",
            data=pdf_bytes,
            file_name=pdf_filename,
            mime="application/pdf",
            use_container_width=True,
        )


def render_assistant(event: dict) -> None:
    st.markdown('<div class="section-title"><span>→</span>AI Procurement Assistant</div>', unsafe_allow_html=True)

    approval_ready = st.session_state.approval_ready
    if approval_ready:
        st.markdown(
            '<div class="approval-box approval-ready">후보를 선택하고 발주 문서를 생성할 수 있습니다</div>',
            unsafe_allow_html=True,
        )
    else:
        st.markdown(
            '<div class="approval-box approval-locked">대체 후보 검색을 먼저 실행하세요</div>',
            unsafe_allow_html=True,
        )

    with st.container(height=440):
        for message in st.session_state.messages:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

    prompt = st.chat_input("재고 확인해줘 / 대체품 찾아줘 / 승인 / 반려")
    if prompt:
        st.session_state.messages.append({"role": "user", "content": prompt})
        response = assistant_reply(prompt, event)
        st.session_state.messages.append({"role": "assistant", "content": response})
        st.rerun()

    col1, col2 = st.columns(2)
    if col1.button("상태 초기화", use_container_width=True):
        reset_demo_state()
        st.rerun()
    if col2.button("검색 실행", use_container_width=True):
        response = assistant_reply("대체품 찾아줘", event)
        st.session_state.messages.append({"role": "assistant", "content": response})
        st.rerun()



# ── Main ───────────────────────────────────────────────────────────────────────

def main() -> None:
    st.markdown(PROFESSIONAL_CSS, unsafe_allow_html=True)
    initialize_session()

    inventory = load_inventory_from_db()
    shortages = shortage_rows(inventory)
    has_serpapi = bool(os.environ.get("SERPAPI_API_KEY"))

    render_page_header(has_serpapi)

    if shortages.empty:
        st.success("현재 부족 자재가 없습니다. 모든 재고가 안전 수준 이상입니다.")
        return

    selected_default = default_selected_material(shortages)

    nav, workspace = st.columns([1.25, 5], gap="large")
    with nav:
        selected_material = render_sidebar(shortages, selected_default)

    st.session_state.selected_material_id = selected_material
    clear_pipeline_state_for_material_change(selected_material)

    events = shortage_events_by_material()
    event = events[selected_material]

    with workspace:
        render_step_tracker(st.session_state.step, st.session_state.approval_ready)
        render_top_metrics(event)

        left, right = st.columns([7, 3.2], gap="large")
        with left:
            render_inventory(inventory, shortages)
            if st.session_state.get("candidate_results") or st.session_state.get("last_error"):
                st.divider()
                render_pipeline_results(event)
            if st.session_state.get("po_created") and PO_FILE.exists():
                st.divider()
                render_po(event)
        with right:
            render_assistant(event)


if __name__ == "__main__":
    main()
